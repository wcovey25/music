"""sheet.py — a CSV/XLSX spreadsheet of songs (V1/V2 format) as a Collection."""
import csv
import math
import os

from ..core.models import Collection, Track
from .base import ResolveError


def _rows(path):
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ResolveError("Reading .xlsx files needs one extra package:  pip install openpyxl") from None
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            it = wb.active.iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it, [])]
            for row in it:
                yield dict(zip(head, row))
        finally:
            wb.close()
    else:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            yield from csv.DictReader(fh)


def _clean(row):
    out = {}
    for k, v in row.items():
        if k is None or v is None or (isinstance(v, float) and math.isnan(v)):
            continue
        out[str(k).strip().lower()] = v.strip() if isinstance(v, str) else str(v)
    return out


def load(path):
    """Songs from a spreadsheet with 'track' and 'artist' columns (plus optional album / mb_* columns)."""
    tracks = []
    try:
        for raw in _rows(path):
            row = _clean(raw)
            title = row.get("track") or row.get("title") or row.get("mb_title") or ""
            artist = row.get("artist") or row.get("mb_artist") or ""
            if not title:
                continue
            year = (row.get("mb_date") or row.get("year") or "")[:4]
            tracks.append(Track(title=title, artist=artist or "Unknown", album=row.get("album") or row.get("mb_release") or "",
                                year=year if year.isdigit() else "", mbid=(row.get("mb_recording_id") or "").strip(),
                                service="sheet"))
    except (OSError, UnicodeDecodeError) as e:
        raise ResolveError(f"Couldn’t read {os.path.basename(path)} ({e})") from None
    if not tracks:
        raise ResolveError("No songs found. The spreadsheet needs a 'track' and an 'artist' column.")
    return Collection(title=os.path.splitext(os.path.basename(path))[0], tracks=tracks, service="sheet", kind="sheet",
                      source=path)
